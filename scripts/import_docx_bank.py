"""Extract the supplied 11 chapter documents into a deployable, traceable bank."""
import argparse
import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
CHAPTERS = ["网络体系结构", "数据通信基础", "网络层协议", "传输层协议", "路由协议", "局域网与交换技术", "IP地址与子网划分", "网络安全", "网络管理", "应用层协议", "网络操作系统与服务器配置"]
NUMBERS = ["一", "二", "三", "四", "五", "六", "七", "八", "九", "十", "十一"]
# Ordered stem-based classification; chapter fallback is explicit, never invented.
TOPICS = [
    (r"DR|BDR", "OSPF DR/BDR"), (r"Router ID", "OSPF Router ID"),
    (r"OSPF.*(?:邻居|邻接)", "OSPF 邻居状态"), (r"OSPF", "OSPF"),
    (r"BGP", "BGP"), (r"RIP", "RIP"), (r"IS.IS", "IS-IS"),
    (r"静态路由|默认路由|路由优先级|管理距离|最长.*匹配", "路由选择"),
    (r"广播地址", "子网广播地址"), (r"网络地址|网络号", "子网网络地址"),
    (r"主机.*(?:数|地址)|子网.*(?:数|划分)|掩码|CIDR|VLSM|路由汇总", "子网划分与 CIDR"),
    (r"NDP|邻居发现", "IPv6 邻居发现"), (r"IPv6", "IPv6"), (r"ARP", "ARP"),
    (r"ICMP|ping|tracert|traceroute", "ICMP"), (r"NAT|NAPT|PAT", "NAT"),
    (r"IP.*地址|IPv4|IP.*首部|IP.*分片|TTL", "IPv4"),
    (r"VLAN|Trunk|Access|802\.1Q", "VLAN"), (r"STP|生成树|根桥|BPDU", "STP"),
    (r"以太网|MAC|CSMA|交换机|双工|冲突域|广播域", "以太网与交换"),
    (r"WLAN|无线|802\.11|SSID|WPA", "无线网络"),
    (r"TCP.*(?:握手|连接|SYN|FIN)", "TCP 连接"), (r"TCP", "TCP"), (r"UDP", "UDP"),
    (r"OSI|TCP/IP.*(?:模型|体系)|PDU|封装|服务访问点", "OSI 与 TCP/IP 模型"),
    (r"奈奎斯特|奈氏|Nyquist", "奈奎斯特定理"), (r"香农|信噪比", "香农定理"),
    (r"编码|码元|Baud|调制|采样|量化|PCM", "编码与调制"),
    (r"复用|WDM|TDM|FDM|CDM", "多路复用"), (r"CRC|海明|校验|误码", "差错控制"),
    (r"光纤|双绞线|同轴|传输介质", "传输介质"),
    (r"ACL|访问控制列表", "ACL"), (r"防火墙|WAF|IDS|IPS", "防火墙与入侵防御"),
    (r"加密|RSA|AES|DES|SM\d|密码|密钥|数字签名|证书|PKI|哈希|散列", "密码与身份认证"),
    (r"VPN|IPSec|IPsec|SSL|TLS", "VPN 与安全协议"),
    (r"攻击|病毒|木马|漏洞|DDoS|钓鱼", "网络攻击与防护"),
    (r"SNMP|MIB|RMON", "SNMP 与网络监控"), (r"Syslog|日志", "日志管理"),
    (r"DNS|域名", "DNS"), (r"DHCP", "DHCP"), (r"FTP", "FTP"),
    (r"SMTP|POP3|IMAP|邮件", "电子邮件"), (r"HTTP|WWW|Web|WEB|CDN", "Web 与 HTTP"),
    (r"NTP", "NTP"), (r"SSH|Telnet", "远程管理"),
    (r"Linux.*权限|chmod|chown|chgrp", "Linux 权限"), (r"Linux|Shell|Bash", "Linux 命令与管理"),
    (r"Windows|活动目录|IIS|组策略", "Windows 服务器"),
]

def normalize(text):
    return re.sub(r"[\s，。？?！!、：:]", "", text).lower()

def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    args = parser.parse_args()
    records, documents = [], []
    for chapter_no, (cn, chapter) in enumerate(zip(NUMBERS, CHAPTERS), 1):
        name = f"第{cn}章 {chapter}（含答案解析）.docx"
        file = args.source / name
        with ZipFile(file) as archive:
            if any(n.startswith("word/media/") for n in archive.namelist()):
                raise ValueError(f"Embedded media needs manual extraction: {name}")
        doc = Document(file)
        if doc.tables or doc.element.xpath(".//m:oMath"):
            raise ValueError(f"Tables or equations need review: {name}")
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        declared = re.search(r"本章共\s*(\d+)\s*道", "\n".join(paragraphs[:3]))
        questions, current = [], None
        for line in paragraphs[3:]:
            if line.startswith("—— 本章完"):
                continue
            start = re.match(r"^(\d+)\.\s*(.+)$", line)
            if start:
                if current:
                    questions.append(current)
                current = {"number": int(start[1]), "question": start[2], "options": {}, "answer": [], "analysis": ""}
            elif current:
                option = re.match(r"^([ABCD])\.\s*(.+)$", line)
                if option:
                    current["options"][option[1]] = option[2]
                elif line.startswith("【答案】"):
                    answer = line.removeprefix("【答案】").strip()
                    if not re.fullmatch(r"[ABCD、,，\s]+", answer):
                        raise ValueError(f"Unrecognized answer {name}:{current['number']}")
                    current["answer"] = re.findall(r"[ABCD]", answer)
                elif line.startswith("【解析】"):
                    current["analysis"] = line.removeprefix("【解析】")
                else:
                    raise ValueError(f"Unparsed paragraph {name}: {line}")
        if current:
            questions.append(current)
        assert declared and len(questions) == int(declared[1]) == 50, name
        assert [q["number"] for q in questions] == list(range(1, 51)), name
        digest = hashlib.sha256(file.read_bytes()).hexdigest()
        documents.append({"file": name, "sha256": digest, "questionCount": len(questions), "images": 0, "tables": 0})
        for raw in questions:
            assert set(raw["options"]) == set("ABCD") and raw["analysis"] and raw["answer"], (name, raw)
            assert len(set(raw["options"].values())) == 4, (name, raw["number"])
            topic = next((topic for pattern, topic in TOPICS if re.search(pattern, raw["question"], re.I)), chapter)
            number = raw.pop("number")
            records.append({
                "id": f"collection-c{chapter_no:02d}-q{number:03d}",
                **raw,
                "type": "single_choice" if len(raw["answer"]) == 1 else "multiple_choice",
                "chapter": "IPv4 与子网划分" if chapter_no == 7 else chapter,
                "knowledgePoint": topic,
                "difficulty": "medium",
                "tags": [topic, "软考真题汇编"],
                "provenance": [{"file": name, "questionNumber": number, "sha256": digest}],
                "sourceVerification": "用户提供汇编，未核实官方出处与考试年份",
            })
    unique, duplicates, review = {}, [], []
    for q in records:
        key = normalize(q["question"]) + json.dumps(q["options"], ensure_ascii=False, sort_keys=True)
        if key in unique:
            original = unique[key]
            same = q["options"] == original["options"] and q["answer"] == original["answer"]
            if same:
                original["provenance"].extend(q["provenance"])
                duplicates.append({"id": q["id"], "retainedId": original["id"], "analysis": q["analysis"], "provenance": q["provenance"]})
            else:
                review.append({"reason": "重复题干的选项或答案不一致", "question": q})
        else:
            unique[key] = q
    # The source itself contradicts this answer: the explanation says OSI appeared later.
    for key, q in list(unique.items()):
        if q["id"] == "collection-c01-q050":
            review.append({"reason": "题目要求选错误说法，但 D 项与解析均称 OSI 出现较晚；原答案 D 不成立，暂不发布", "question": q})
            del unique[key]
    destination = ROOT / "server" / "question-banks" / "network-engineer"
    question_dir = destination / "questions" / "user-collection"
    report_dir = destination / "reports"
    question_dir.mkdir(parents=True, exist_ok=True)
    report_dir.mkdir(parents=True, exist_ok=True)
    bank = list(unique.values())
    report = {"sourceCount": len(records), "publishedCount": len(bank), "duplicateCount": len(duplicates), "reviewCount": len(review), "documents": documents, "duplicates": duplicates, "needsReview": review}
    for index, chapter in enumerate(CHAPTERS, 1):
        stored_chapter = "IPv4 与子网划分" if index == 7 else chapter
        rows = [q for q in bank if q["chapter"] == stored_chapter]
        (question_dir / f"{index:02d}.json").write_text(json.dumps(rows, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (report_dir / "user-collection.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: report[k] for k in ["sourceCount", "publishedCount", "duplicateCount", "reviewCount"]}, ensure_ascii=False))

if __name__ == "__main__":
    main()
